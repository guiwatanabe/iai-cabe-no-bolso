from pptx import Presentation
P = '/Users/manasses/Documents/ChatGPT/Batalha de Agentes/output/'
prs = Presentation(P + 'canvas_neo_time05.pptx')
n = 0
for sh in prs.slides[0].shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if 'NEO' in r.text:
                    print('antes:', repr(r.text)); r.text = r.text.replace('NEO', 'Cabe no Bolso'); n += 1; print('depois:', repr(r.text))
assert n >= 1, 'NEO nao encontrado'
out = P + 'canvas_cabe_no_bolso_time05.pptx'; prs.save(out)
print('OK', out, '|', ' | '.join(sh.text_frame.text.replace('\n', ' / ') for sh in Presentation(out).slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip())[:400])
