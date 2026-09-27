import pandas as pd, numpy as np
from pptx import Presentation
P = '/Users/manasses/Documents/ChatGPT/Batalha de Agentes/'
df = pd.read_csv(P + 'data/extrato_sintetico.csv.gz', dtype={'id_usuario': str, 'descr': str})
df['descr'] = df['descr'].fillna(''); df['c'] = (df.vlr * 100).round().astype('int64')
df['card'] = df.descr.str.startswith('cart credito'); df['isfat'] = df.nom_cate_micro == 'Pagamento de fatura'
uid = '0a37f67b-77d5-4ab8-af28-2d88aaa252f1'; u = df[df.id_usuario == uid]
e = u[u.tipo == 'E'].groupby('nom_cate_micro').agg(n=('c', 'size'), med=('c', 'median'), soma=('c', 'sum'))
print('ENTRADAS persona'); print((e / [1, 100, 100]).round(2).to_string())
s = u[(u.tipo == 'S') & ~u.card].groupby('nom_cate_micro').agg(n=('c', 'size'), meses=('anomes', 'nunique'), med=('c', 'median'), soma=('c', 'sum'))
s = s[s.meses >= 10].sort_values('soma', ascending=False)
print('SAIDAS recorrentes (>=10 meses, sem compras no cartao)'); print((s / [1, 1, 100, 100]).round(2).to_string())
g = df.assign(E=np.where(df.tipo == 'E', df.c, 0), S=np.where((df.tipo == 'S') & ~df.card & ~df.isfat, df.c, 0)).groupby(['id_usuario', 'anomes'])[['E', 'S']].sum()
g['sobra'] = (g.E - g.S) / 100
fat = df[df.isfat].copy(); fat['parc'] = ~fat.descr.str.contains('integral')
fat = fat.set_index(['id_usuario', 'anomes']).join(g['sobra'])
pay = set(fat[fat.parc].index.get_level_values(0)); f3 = fat[fat.index.get_level_values(0).isin(pay)]
print('TESTE sobra do mes (entradas - saidas sem cartao e sem fatura), 689 pagadores: parcial', round(f3[f3.parc].sobra.median(), 2), '| integral', round(f3[~f3.parc].sobra.median(), 2))
print('valor pago mediano: parcial', f3[f3.parc].c.median() / 100, '| integral', f3[~f3.parc].c.median() / 100)
extra = set(map(tuple, df[df.nom_cate_micro.isin(['13o salario', 'Bonus PLR'])][['id_usuario', 'anomes']].values))
f3 = f3.assign(extra=[i in extra for i in f3.index])
print('taxa parcial com renda extra no mes:', f3.groupby('extra').parc.agg(['mean', 'size']).round(3).to_dict())
print('persona sobra por mes:', g.loc[uid].sobra.round(0).to_dict())
prs = Presentation(P + 'output/ficha_cabe_no_mes_time05_rascunho.pptx')
for sh in prs.slides[0].shapes:
    if sh.has_text_frame:
        for i, p in enumerate(sh.text_frame.paragraphs):
            print('PPTX', sh.name, '| p', i, '|', [(r.text, r.font.bold, r.font.size.pt if r.font.size else None) for r in p.runs])
