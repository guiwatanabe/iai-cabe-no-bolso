import pandas as pd, numpy as np
from pptx import Presentation
P = '/Users/manasses/Documents/ChatGPT/Batalha de Agentes/'
df = pd.read_csv(P + 'data/extrato_sintetico.csv.gz', dtype={'id_usuario': str, 'descr': str})
df['descr'] = df['descr'].fillna(''); df['c'] = (df.vlr * 100).round().astype('int64')
df['card'] = df.descr.str.startswith('cart credito'); df['isfat'] = df.nom_cate_micro == 'Pagamento de fatura'
fat = df[df.isfat].copy(); fat['parc'] = ~fat.descr.str.contains('integral'); fat = fat.set_index(['id_usuario', 'anomes'])
for nome, emask in [('com PIX recebido', df.tipo == 'E'), ('sem Recebimentos diversos', (df.tipo == 'E') & (df.nom_cate_macro != 'Recebimentos diversos'))]:
    g = df.assign(E=np.where(emask, df.c, 0), S=np.where((df.tipo == 'S') & ~df.card & ~df.isfat, df.c, 0)).groupby(['id_usuario', 'anomes'])[['E', 'S']].sum()
    sob = (g.E - g.S) / 100
    med_int = fat[~fat.parc].groupby(level=0).c.median() / 100
    fp = fat[fat.parc & fat.index.get_level_values(0).isin(med_int.index)]
    ok = sob.reindex(fp.index).values >= med_int.reindex(fp.index.get_level_values(0)).values
    print(nome, '| meses parciais c/ sobra >= fatura integral tipica:', round(ok.mean() * 100, 1), '% de', len(fp))
T = {
 'Retângulo: Cantos Arredondados 3': {0: ['Nome do Agente: Farol da Fatura']},
 'Retângulo: Cantos Arredondados 4': {1: ['Cliente paga parte da fatura sem ver o custo.'], 3: ['Na base sintética: ', '689 de 1.000 pagaram parcial; juros em 99,97% desses meses.']},
 'Retângulo: Cantos Arredondados 5': {1: ['Cliente CLT decide quanto pagar da fatura.'], 3: ['Contexto: ', 'vencimento em 3 dias e contas fixas a pagar.']},
 'Retângulo: Cantos Arredondados 6': {1: ['Mostra o custo de cada forma de pagar a fatura.'], 3: ['Três capacidades: ', 'Lê o mês; compara custos; orienta sem vender crédito.']},
 'Retângulo: Cantos Arredondados 8': {0: ['Dados e tecnologia:', 'BigQuery → cálculo de custo → Gemini → Cloud Run.'], 1: ['Fontes: ', 'Extrato + taxas como parâmetro explícito.']},
 'Retângulo: Cantos Arredondados 9': {0: ['Como mediremos valor:', 'Menos faturas pagas no mínimo; juros evitados.'], 1: ['Sinais no piloto: ', '% integral ou parcelada; R$ de juros evitados; compreensão.']},
 'Retângulo: Cantos Arredondados 10': {0: ['Escopo da Demo', 'Fatura dia 15: alerta → 3 opções → plano escolhido.'], 1: ['Limites: ', 'Sem pagamento real, taxas ilustrativas, dados de 2025.']},
}
prs = Presentation(P + 'output/ficha_cabe_no_mes_time05_rascunho.pptx')
for sh in prs.slides[0].shapes:
    if sh.name in T:
        for i, texts in T[sh.name].items():
            runs = sh.text_frame.paragraphs[i].runs
            assert len(runs) == len(texts), (sh.name, i, len(runs))
            for r, t in zip(runs, texts): r.text = t
out = P + 'output/ficha_farol_da_fatura_time05_v2.pptx'; prs.save(out)
for sh in Presentation(out).slides[0].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip(): print('OK', sh.text_frame.text.replace('\n', ' | '))
